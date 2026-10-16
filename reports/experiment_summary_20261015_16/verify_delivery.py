"""Read-only verification of exported XLSX/CSV, then package deliverables."""
import csv
import hashlib
import json
import math
from pathlib import Path
import zipfile
import openpyxl

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
OUT=ROOT/'outputs/01a11c24-c60c-7b42-99e6-70e3c00dea5e/experiment_summary_20261015_16'
d=json.loads((HERE/'summary_data.json').read_text(encoding='utf-8'))
sources=json.loads((HERE/'source_manifest.json').read_text(encoding='utf-8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
for s in sources:
    assert sha(ROOT/s['path'])==s['sha256'], f"Source changed: {s['id']}"
assert (OUT/'summary_data.json').read_bytes()==(HERE/'summary_data.json').read_bytes()
with (OUT/'question_results.csv').open(encoding='utf-8-sig',newline='') as f:
    rows=list(csv.DictReader(f))
assert len(rows)==7148
assert len({(r['batch'],r['split'],r['question_id'],r['arm'],r['density']) for r in rows})==len(rows)
assert sum(r['split']=='test' for r in rows)==6332
assert sum(r['split']=='val' for r in rows)==816
assert sum(r['generation_status']=='invalid' for r in rows)==9
for x in d['test_arms']:
    rr=[r for r in rows if r['split']=='test' and r['dataset']==x['dataset'] and r['arm']==x['arm'] and float(r['density'])==x['density']]
    assert len(rr)==x['questions']
    assert sum(int(r['correct_or_passed']) for r in rr)==x['correct']
    assert sum(int(r['logical_model_tokens']) for r in rr)==x['logical_tokens']

xlsx=OUT/'实验数据汇总_20261015-16.xlsx'
values=openpyxl.load_workbook(xlsx,data_only=True)
formulas=openpyxl.load_workbook(xlsx,data_only=False)
assert len(values.sheetnames)==6
ws=values['Test对照']
for i,x in enumerate(d['test_comparisons'],7):
    assert ws.cell(i,1).value==x['dataset']
    expected={4:x['questions'],5:x['moc_correct'],6:x['candidate_correct'],7:x['moc_correct']/x['questions'],8:x['candidate_correct']/x['questions'],9:(x['candidate_correct']-x['moc_correct'])/x['questions']*100,10:x['ci_low']*100,11:x['ci_high']*100,12:x['mcnemar_p'],13:1-x['candidate_tokens']/x['moc_tokens'],14:x['moc_tokens'],15:x['candidate_tokens']}
    for c,v in expected.items():
        assert math.isclose(ws.cell(i,c).value,v,rel_tol=1e-10,abs_tol=1e-10),(i,c)
    assert formulas['Test对照'].cell(i,7).value==f'=E{i}/D{i}'
assert [values['批次成本'].cell(15,i).value for i in range(7,12)]==[39181781,69536,0,39251317,60218]
for i,x in enumerate(d['val_arms'],7):
    assert values['Val对照'].cell(i,5).value==x['correct']
    assert math.isclose(values['Val对照'].cell(i,6).value,x['correct']/x['questions'])
for i,x in enumerate(d['routing'],7):
    assert math.isclose(values['路由诊断'].cell(i,8).value,x['empty']/x['nonroot_decisions'])
for i,x in enumerate(d['test_arms'],7):
    assert values['Test用量明细'].cell(i,2).value==('原MOC' if x['arm']=='moc' else '任务适配' if x['arm']=='adapted' else 'v2.1')
    assert values['Test用量明细'].cell(i,8).value==x['logical_tokens']
errors=[]
for sheet in values:
    for row in sheet:
        for c in row:
            if c.data_type=='e': errors.append((sheet.title,c.coordinate,c.value))
assert not errors, errors
for sheet in formulas:
    for tbl in sheet.tables.values():
        assert tbl.autoFilter is not None, (sheet.title,tbl.name)
assert 'matched 0 entries' in (OUT/'workbook_formula_scan.ndjson').read_text(encoding='utf-8')
assert len(list(OUT.glob('preview_*.png')))==8

verification={'status':'passed','source_files_sha256_verified':len(sources),'test_comparisons':12,'test_arms':24,'question_rows':7148,'test_rows':6332,'val_rows':816,'unique_condition_keys':True,'correct_counts_and_logical_tokens_match_raw_extract':True,'xlsx_exported_cached_values_match':True,'xlsx_formula_cells_present':True,'xlsx_error_cells':errors,'xlsx_filters_present':True,'input_change_recalculation':'verified in artifact-tool before export; audited inputs restored','visual_qa':'all six worksheets rendered; eight ranges inspected; labels, wrapping and arm names corrected','new_model_calls':0,'totals':d['totals'],'limits':['local artifact consistency audit, not independent reproduction','bootstrap intervals retained from audited original analyses','logical tokens are known main+merger tokens including cache hits','seven-batch cumulative cost, not strict calendar-day spending']}
for p in [HERE/'verification.json', OUT/'verification.json']:
    p.write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf-8')

readme='''MOC实验汇总：2026-10-15—16（北京时间）

先读《实验数据汇总_20261015-16.md》，再用Excel筛选。CSV为UTF-8 BOM；JSON为UTF-8。
逐题表7,148行，test 6,332行、val 816行；同一题的不同密度/方法是配对条件，不是独立样本。
test_comparisons.csv的training是候选v2.1策略的训练方式，原MOC没有本轮适配训练。
正确数、逻辑token、actual/unknown/pending分列；费用是七批累计，不是严格日历两天消费。
原始提示词、测试金标、生成代码、大型raw归档不在此整理包；来源路径及SHA在source_manifest.json。
本地一致性核验见verification.json；文件字节哈希见SHA256SUMS.json。不宣称独立复现。
'''
(OUT/'README.txt').write_text(readme,encoding='utf-8')
names=['实验数据汇总_20261015-16.xlsx','实验数据汇总_20261015-16.md','question_results.csv','test_comparisons.csv','val_arms.csv','cost_ledger.csv','routing.csv','summary_data.json','source_manifest.json','verification.json','README.txt']
manifest=[{'file':n,'bytes':(OUT/n).stat().st_size,'sha256':sha(OUT/n)} for n in names]
(OUT/'SHA256SUMS.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
bundle=OUT/'实验数据整理包_20261015-16.zip'
with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
    for n in names+['SHA256SUMS.json']:z.write(OUT/n,n)
with zipfile.ZipFile(bundle) as z:
    assert z.testzip() is None
    for item in manifest:assert hashlib.sha256(z.read(item['file'])).hexdigest()==item['sha256']
print(json.dumps({'status':'passed','sources':len(sources),'rows':len(rows),'sheets':values.sheetnames,'bundle':str(bundle),'bundle_bytes':bundle.stat().st_size,'bundle_sha256':sha(bundle)},ensure_ascii=True))
